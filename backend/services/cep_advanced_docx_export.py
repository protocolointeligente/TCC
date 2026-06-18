from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json


def _set_margins(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Cm(3)
        section.left_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2)


def _set_font(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


def _heading(document: Document, text: str, level: int = 1) -> None:
    h = document.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _para(document: Document, text: str, indent: bool = True) -> None:
    if not text or not text.strip():
        return
    p = document.add_paragraph(text.strip())
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _severity_label(severity: str) -> str:
    return {"error": "[ERRO]", "warning": "[ALERTA]", "info": "[INFO]"}.get(severity, severity.upper())


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        run = hdr_cells[i].paragraphs[0].runs[0] if hdr_cells[i].paragraphs[0].runs else hdr_cells[i].paragraphs[0].add_run(h)
        run.bold = True
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val) if val else "—"
    document.add_paragraph()


def export_advanced_cep_to_docx(project, result: dict, filename: str) -> str:
    doc = Document()
    _set_margins(doc)
    _set_font(doc)

    # ── Cover ──────────────────────────────────────────────────────────────
    title_para = doc.add_heading(project.title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    meta_lines = [
        f"Curso: {project.course or '—'}",
        f"Área: {project.area or '—'}",
        f"Tipo de estudo: {project.study_type or '—'}",
        f"Instituição: {project.institution or '—'}",
        f"Orientador: {project.advisor or '—'}",
        f"População: {project.target_population or '—'} | Faixa etária: {project.age_group or '—'}",
    ]
    for line in meta_lines:
        _para(doc, line, indent=False)
    doc.add_page_break()

    # ── 1. Validador ético ─────────────────────────────────────────────────
    _heading(doc, "1. Validador Ético")
    risk = result.get("risk", {})
    _para(doc, f"Nível de risco: {risk.get('risk_level', '—')} (score: {risk.get('risk_score', '—')})", indent=False)
    for r in risk.get("identified_risks", []):
        _para(doc, f"• {r}", indent=False)
    doc.add_paragraph()
    issues = result.get("validator", {}).get("issues", [])
    if issues:
        _heading(doc, "Pendências do validador", level=2)
        for issue in issues:
            _para(doc, f"{_severity_label(issue['severity'])} [{issue['field']}] {issue['message']}", indent=False)
    else:
        _para(doc, "Nenhum problema crítico detectado pelo validador automático.", indent=False)
    doc.add_paragraph()

    # ── 2. Resoluções éticas ───────────────────────────────────────────────
    resolution_checks = result.get("ethics_resolution_checks", [])
    if resolution_checks:
        _heading(doc, "2. Resoluções Éticas Aplicáveis")
        for c in resolution_checks:
            _para(doc, f"{c['resolution']} — Status: {c['status']}", indent=False)
            _para(doc, f"  Recomendação: {c['recommendation']}", indent=False)
        doc.add_paragraph()

    # ── 3. Coerência metodológica ──────────────────────────────────────────
    coherence = result.get("methodological_coherence", {})
    _heading(doc, "3. Coerência Metodológica")
    _para(doc, f"Status: {coherence.get('status', '—')}", indent=False)
    for prob in coherence.get("problems", []):
        _para(doc, f"• {prob}", indent=False)
    if coherence.get("recommendation"):
        _para(doc, coherence["recommendation"], indent=False)
    doc.add_paragraph()

    # ── 4. Matriz metodológica ─────────────────────────────────────────────
    matrix = result.get("objective_matrix", [])
    if matrix:
        _heading(doc, "4. Matriz Objetivo → Instrumento → Variável → Análise")
        _add_table(
            doc,
            ["Objetivo específico", "Instrumento", "Variável", "Análise", "Desfecho"],
            [[r.get("objective",""), r.get("instrument",""), r.get("variable",""), r.get("analysis",""), r.get("outcome","")] for r in matrix],
        )

    # ── 5. Instrumentos sugeridos ──────────────────────────────────────────
    instruments = result.get("instruments", [])
    if instruments:
        _heading(doc, "5. Instrumentos de Pesquisa Sugeridos")
        for inst in instruments:
            _para(doc, f"{inst.get('name','?')} — {inst.get('construct','')}", indent=False)
            if inst.get("validated_population_brazil"):
                _para(doc, f"  Validação brasileira: {inst['validated_population_brazil']}", indent=False)
            if inst.get("license_use"):
                _para(doc, f"  Licença/uso: {inst['license_use']}", indent=False)
        doc.add_paragraph()

    # ── 6. Advisor estatístico ────────────────────────────────────────────
    stats = result.get("statistics", {})
    if stats.get("suggestions"):
        _heading(doc, "6. Análises Estatísticas Sugeridas")
        for s in stats["suggestions"]:
            _para(doc, f"• {s['test']}: {s['use']}", indent=False)
        if stats.get("software"):
            _para(doc, f"Softwares recomendados: {', '.join(stats['software'])}", indent=False)
        doc.add_paragraph()

    # ── 7. TCLE ───────────────────────────────────────────────────────────
    doc.add_page_break()
    smart_tcle = result.get("smart_tcle") or (project.smart_tcle if hasattr(project, "smart_tcle") else None)
    if smart_tcle:
        _heading(doc, "7. Termo de Consentimento Livre e Esclarecido (TCLE)")
        for line in smart_tcle.split("\n"):
            if line.startswith("## "):
                _heading(doc, line[3:].strip(), level=2)
            elif line.strip():
                _para(doc, line.strip())
        doc.add_paragraph()

    # ── 8. Checklist Plataforma Brasil ────────────────────────────────────
    plataforma = result.get("plataforma_brasil", {})
    checklist_sections = plataforma.get("checklist_sections", [])
    if checklist_sections:
        doc.add_page_break()
        _heading(doc, "8. Checklist Plataforma Brasil")
        for section in checklist_sections:
            _heading(doc, section["section"], level=2)
            for field in section["fields"]:
                _para(doc, f"☐  {field}", indent=False)
        if plataforma.get("plataforma_brasil_tips"):
            _heading(doc, "Dicas de preenchimento", level=2)
            _para(doc, plataforma["plataforma_brasil_tips"])
        doc.add_paragraph()

    # ── 9. Cronograma ─────────────────────────────────────────────────────
    schedule = result.get("schedule", [])
    if schedule:
        doc.add_page_break()
        _heading(doc, "9. Cronograma de Execução")
        _add_table(
            doc,
            ["Fase", "Mês início", "Mês fim", "Duração"],
            [[r["phase"], str(r["start_month"]), str(r["end_month"]), f"{r['duration_months']} mês(es)"] for r in schedule],
        )

    # ── 10. Orçamento ─────────────────────────────────────────────────────
    budget = result.get("budget", [])
    if budget:
        _heading(doc, "10. Orçamento Estimado")
        _add_table(doc, ["Item", "Custo estimado"], [[b["item"], b["estimated_cost"]] for b in budget])

    # ── 11. Anexos necessários ────────────────────────────────────────────
    attachments = result.get("attachments", {})
    required_docs = attachments.get("required", [])
    if required_docs:
        _heading(doc, "11. Documentos Obrigatórios para Submissão")
        for a in required_docs:
            _para(doc, f"• {a['name']}: {a['description']}", indent=False)
        doc.add_paragraph()

    # ── 12. Pendências simuladas ──────────────────────────────────────────
    pendencies = result.get("cep_pendency_simulator", [])
    if pendencies:
        doc.add_page_break()
        _heading(doc, "12. Simulação de Pendências do CEP")
        for p in pendencies:
            _heading(doc, f"{p['type']} — probabilidade {p['probability']}", level=2)
            _para(doc, f"Motivo: {p['reason']}", indent=False)
            _para(doc, f"Comentário esperado do CEP: {p['cep_possible_comment']}", indent=False)
            _para(doc, f"Como corrigir: {p['how_to_fix']}", indent=False)
        doc.add_paragraph()

    # ── 13. Pacote pré-submissão ──────────────────────────────────────────
    pre_sub = result.get("pre_submission_package", {})
    if pre_sub:
        _heading(doc, "13. Pacote de Pré-Submissão")
        _para(doc, f"Status: {pre_sub.get('status','—')}", indent=False)
        for doc_name in pre_sub.get("required_documents", []):
            _para(doc, f"☐  {doc_name}", indent=False)
        for missing in pre_sub.get("missing_or_pending", []):
            _para(doc, f"⚠  {missing}", indent=False)
        _para(doc, pre_sub.get("final_instruction", ""), indent=False)
        doc.add_paragraph()

    # ── 14. Parecer do orientador ─────────────────────────────────────────
    orientador = result.get("orientador_opinion") or (project.orientador_opinion if hasattr(project, "orientador_opinion") else None)
    if orientador:
        doc.add_page_break()
        _heading(doc, "14. Parecer do Orientador (gerado por IA)")
        _para(doc, orientador)

    doc.save(filename)
    return filename
