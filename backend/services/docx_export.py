from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from services.abnt import generate_abnt_references


def export_review_to_docx(
    title: str,
    review_text: str,
    articles: list[dict],
    filename: str,
) -> str:
    document = Document()

    for section in document.sections:
        section.top_margin = Cm(3)
        section.left_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    heading = document.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in review_text.split("\n"):
        if block.strip():
            p = document.add_paragraph(block.strip())
            p.paragraph_format.first_line_indent = Cm(1.25)
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.add_page_break()
    document.add_heading("Referências", level=1)
    for ref in generate_abnt_references(articles):
        p = document.add_paragraph(ref)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.save(filename)
    return filename
